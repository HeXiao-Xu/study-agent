"""S1-1 解析模块测试（spec §7）：自造 fixture，零网络零模型。"""
from __future__ import annotations

from pathlib import Path

import docx as python_docx
import pymupdf
import pytest

from app.ai.parse.models import BlockType, ParsedBlock, ParsedDocument


def test_block_defaults() -> None:
    """heading_path 抽不到默认空列表；type 默认 text（不瞎编）。"""
    b = ParsedBlock(text="通项公式", page_no=1)
    assert b.heading_path == []
    assert b.type is BlockType.text


def test_block_docx_page_none() -> None:
    """docx 无分页概念：page_no 允许 None（spec 设计决策 4）。"""
    b = ParsedBlock(text="错位相减法", page_no=None)
    assert b.page_no is None


def test_document_envelope() -> None:
    d = ParsedDocument(source="a.pdf", parser="pymupdf", blocks=[ParsedBlock(text="x", page_no=1)])
    assert d.is_scanned is False
    assert d.degraded_reason is None
    assert d.parser == "pymupdf"
    assert d.blocks[0].type.value == "text"


# ---- Task 3：fixture 工厂 + PyMuPDF 兜底 + is_scanned ----


def make_pdf(path: Path, texts: list[str | None]) -> Path:
    """texts 每页一项；None=空白页（扫描版模拟）；含 \\n 按行插入。

    为什么按行插入：单行超过页宽（~54 个汉字）会被裁掉，导致提取字符数
    意外低于 is_scanned 阈值（实测 60 字只提出 48 字）。
    """
    with pymupdf.open() as d:
        for t in texts:
            page = d.new_page()
            if t:
                for i, line in enumerate(t.split("\n")):
                    page.insert_text((72, 72 + i * 20), line, fontname="china-s")
        d.save(path)
    return path


def test_pymupdf_blocks_and_pages(tmp_path: Path) -> None:
    from app.ai.parse.pdf_pymupdf import PyMuPDFParser

    pdf = make_pdf(tmp_path / "note.pdf", ["第一章 数列", "等差数列通项公式", "错位相减法适用差比混合"])
    doc = PyMuPDFParser().parse(pdf)
    assert doc.parser == "pymupdf"
    assert doc.is_scanned is False
    assert [b.page_no for b in doc.blocks] == [1, 2, 3]  # 1-based（spec 设计决策 5）
    assert all(b.heading_path == [] for b in doc.blocks)  # 兜底不猜层级（决策 2）
    assert all(b.type.value == "text" for b in doc.blocks)


def test_pymupdf_skips_blank_pages(tmp_path: Path) -> None:
    from app.ai.parse.pdf_pymupdf import PyMuPDFParser

    pdf = make_pdf(tmp_path / "mixed.pdf", ["有文字的页", None, "另一页"])
    doc = PyMuPDFParser().parse(pdf)
    assert [b.page_no for b in doc.blocks] == [1, 3]


def test_detect_scanned_blank_pdf(tmp_path: Path) -> None:
    from app.ai.parse.pdf_pymupdf import detect_scanned

    rich_text = "等差数列的通项公式推导需要一定篇幅文字说明。\n第二行继续补充足够的字符数量以免误判。\n第三行再次补充字符，确保总量超过阈值。"
    assert detect_scanned(make_pdf(tmp_path / "blank.pdf", [None, None, None])) is True
    assert detect_scanned(make_pdf(tmp_path / "text.pdf", [rich_text])) is False


def test_pymupdf_sanitize_control_chars() -> None:
    """验收实测：字体 ToUnicode 缺失会让 ⊆ 提取成 \\x00，必须清掉但保留换行。"""
    from app.ai.parse.pdf_pymupdf import _sanitize

    assert _sanitize("区间 D\x00I：") == "区间 DI："
    assert _sanitize("第一行\n第二行") == "第一行\n第二行"


# ---- Task 4：docx 解析 ----


def make_docx(path: Path) -> Path:
    """标题 1 + 标题 2 + 正文 + 2x2 表格 的最小样张。"""
    d = python_docx.Document()
    d.add_paragraph("数列笔记", style="Heading 1")
    d.add_paragraph("错位相减法", style="Heading 2")
    d.add_paragraph("适用于等差乘等比型数列求和。")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "方法"
    t.cell(0, 1).text = "适用"
    t.cell(1, 0).text = "错位相减"
    t.cell(1, 1).text = "差比混合"
    d.save(path)
    return path


def test_heading_level_matches_cn_and_en() -> None:
    from app.ai.parse.docx_parser import heading_level

    assert heading_level("Heading 2") == 2
    assert heading_level("标题 3") == 3  # 中文 Word 样式名
    assert heading_level("heading 1") == 1  # 大小写不敏感
    assert heading_level("Normal") is None
    assert heading_level(None) is None


def test_docx_headings_table_and_no_page(tmp_path: Path) -> None:
    from app.ai.parse.docx_parser import DocxParser

    doc = DocxParser().parse(make_docx(tmp_path / "note.docx"))
    assert doc.parser == "docx"
    assert all(b.page_no is None for b in doc.blocks)  # spec 设计决策 4
    assert [b.text for b in doc.blocks[:3]] == ["数列笔记", "错位相减法", "适用于等差乘等比型数列求和。"]
    assert doc.blocks[0].heading_path == []  # h1 自己的路径不含自己
    assert doc.blocks[1].heading_path == ["数列笔记"]
    assert doc.blocks[2].heading_path == ["数列笔记", "错位相减法"]
    table_block = doc.blocks[3]
    assert "| 方法 |" in table_block.text and "| 差比混合 |" in table_block.text
    assert table_block.heading_path == ["数列笔记", "错位相减法"]


# ---- Task 5：Docling 适配器（duck-typed 假对象，不 import docling）----


class _FakeProv:
    def __init__(self, page_no: int) -> None:
        self.page_no = page_no


class _FakeItem:
    def __init__(self, label: str, text: str = "", page_no: int | None = 1) -> None:
        self.label = label
        self.text = text
        self.prov = [_FakeProv(page_no)] if page_no is not None else []


class _FakeTable:
    def __init__(self, page_no: int) -> None:
        self.label = "table"
        self.prov = [_FakeProv(page_no)]

    def export_to_markdown(self) -> str:
        return "| a | b |\n| --- | --- |\n| 1 | 2 |"


class _FakeDoc:
    def __init__(self, items: list) -> None:
        self._items = items

    def iterate_items(self) -> list:
        return self._items


def test_docling_mapping_and_heading_stack() -> None:
    from app.ai.parse.pdf_docling import DoclingParser

    fake = _FakeDoc(
        [
            (_FakeItem("title", "数学笔记"), 0),
            (_FakeItem("section_header", "数列"), 1),
            (_FakeItem("paragraph", "等差数列的通项公式。", page_no=2), None),
            (_FakeItem("page_footer", "第 2 页", page_no=2), None),  # 噪声，必须丢弃
            (_FakeTable(3), None),
        ]
    )
    blocks = DoclingParser()._items_to_blocks(fake)
    assert [b.text for b in blocks] == [
        "数学笔记",
        "数列",
        "等差数列的通项公式。",
        "| a | b |\n| --- | --- |\n| 1 | 2 |",
    ]
    assert blocks[0].heading_path == []  # title 重置栈，自身路径为空
    assert blocks[1].heading_path == ["数学笔记"]
    assert blocks[2].heading_path == ["数学笔记", "数列"]
    assert blocks[3].heading_path == ["数学笔记", "数列"]  # 表格延续当前层级
    assert [b.page_no for b in blocks] == [1, 1, 2, 3]
    assert all(b.type.value == "text" for b in blocks)


def test_docling_non_tuple_iteration() -> None:
    """兼容只 yield 裸 item 的 docling 版本（spec：探针后校正）。"""
    from app.ai.parse.pdf_docling import DoclingParser

    fake = _FakeDoc([_FakeItem("paragraph", "裸 item")])  # 无 tuple 包装
    blocks = DoclingParser()._items_to_blocks(fake)
    assert [b.text for b in blocks] == ["裸 item"]


# ---- Task 6：门面与降级链 ----


def test_facade_unsupported_ext(tmp_path: Path) -> None:
    from app.ai.parse import parse_document
    from app.core.exceptions import InvalidParamError

    f = tmp_path / "a.txt"
    f.write_text("x", encoding="utf-8")
    with pytest.raises(InvalidParamError):
        parse_document(f)


def test_facade_missing_file(tmp_path: Path) -> None:
    from app.ai.parse import parse_document
    from app.core.exceptions import InvalidParamError

    with pytest.raises(InvalidParamError):
        parse_document(tmp_path / "不存在.pdf")


def test_facade_scanned_early_return(tmp_path: Path) -> None:
    from app.ai.parse import parse_document

    doc = parse_document(make_pdf(tmp_path / "blank.pdf", [None, None]))
    assert doc.is_scanned is True
    assert doc.blocks == []
    assert "扫描版" in (doc.degraded_reason or "")


def test_facade_force_pymupdf(tmp_path: Path) -> None:
    from app.ai.parse import parse_document

    rich_text = "快速路径验证：本页提供了足够多的可抽取文字内容。\n第二行补充字符数量避免扫描版误判。\n第三行确保字符总量超过判定阈值。"
    pdf = make_pdf(tmp_path / "note.pdf", [rich_text])
    doc = parse_document(pdf, force_pymupdf=True)
    assert doc.parser == "pymupdf"
    assert doc.degraded_reason and "PyMuPDF" in doc.degraded_reason


def test_facade_docx_dispatch(tmp_path: Path) -> None:
    from app.ai.parse import parse_document

    doc = parse_document(make_docx(tmp_path / "note.docx"))
    assert doc.parser == "docx"


def test_facade_fallback_on_docling_error(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Docling 抛错 → 自动降级 PyMuPDF，degraded_reason 带原因（任务卡要求 5）。"""
    from app.ai.parse import pdf_docling, parse_document

    def boom(self: object, path: Path) -> object:
        raise RuntimeError("docling 内部炸了")

    monkeypatch.setattr(pdf_docling.DoclingParser, "parse", boom)
    rich_text = "降级链验证：本页提供了足够多的可抽取文字内容。\n第二行补充字符数量避免扫描版误判。\n第三行确保字符总量超过判定阈值。"
    pdf = make_pdf(tmp_path / "note.pdf", [rich_text])
    doc = parse_document(pdf)  # force=None 且 Settings 默认 False → 会先试 docling
    assert doc.parser == "pymupdf"
    assert doc.degraded_reason and doc.degraded_reason.startswith("Docling 降级")
    assert doc.blocks and doc.blocks[0].text.startswith("降级链验证")
