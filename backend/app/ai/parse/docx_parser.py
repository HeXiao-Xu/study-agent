"""python-docx 解析器（spec §5.4）。

必须用 iter_inner_content()（python-docx ≥1.1）按文档顺序遍历段落与表格：
document.paragraphs 会丢表格位置，不可用。
"""
from __future__ import annotations

import logging
import re
from pathlib import Path

import docx
import docx.table

from app.ai.parse.base import DocumentParser
from app.ai.parse.models import ParsedBlock, ParsedDocument

logger = logging.getLogger(__name__)

# "Heading 2" / "标题 2" / "heading 2"（中英文 Word 模板都用过）
_HEADING_RE = re.compile(r"^(?:heading|标题)\s*(\d)$", re.IGNORECASE)


def heading_level(style_name: str | None) -> int | None:
    """样式名 → 标题层级数字；非标题返回 None。"""
    if not style_name:
        return None
    m = _HEADING_RE.match(style_name.strip())
    return int(m.group(1)) if m else None


def _table_to_markdown(table: docx.table.Table) -> str:
    rows = [
        "| " + " | ".join(c.text.replace("\n", " ").strip() for c in r.cells) + " |"
        for r in table.rows
    ]
    if not rows:
        return ""
    sep = "| " + " | ".join(["---"] * len(table.rows[0].cells)) + " |"
    return "\n".join([rows[0], sep, *rows[1:]])


class DocxParser(DocumentParser):
    name = "docx"

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() == ".docx"

    def parse(self, path: Path) -> ParsedDocument:
        d = docx.Document(str(path))
        blocks: list[ParsedBlock] = []
        headings: list[str] = []
        skipped_images = 0
        for item in d.iter_inner_content():
            if isinstance(item, docx.table.Table):
                md = _table_to_markdown(item)
                if md:
                    blocks.append(ParsedBlock(text=md, page_no=None, heading_path=list(headings)))
                continue
            # Paragraph：图片以 drawing 形式内嵌，先数后跳（图片识别是下一任务）
            skipped_images += len(item._p.xpath(".//w:drawing"))
            level = heading_level(item.style.name if item.style is not None else None)
            text = item.text.strip()
            if level is not None:
                if text:
                    # 层级 1 重置栈、层级 n 截断到 n-1 再压入：h1→[h1]；h2→[h1,h2]
                    headings = headings[: level - 1] if level > 1 else []
                    headings.append(text)
                    blocks.append(ParsedBlock(text=text, page_no=None, heading_path=list(headings[:-1])))
                continue
            if text:
                blocks.append(ParsedBlock(text=text, page_no=None, heading_path=list(headings)))
        if skipped_images:
            logger.info("docx 跳过 %d 个图片对象（图片识别是下一任务）", skipped_images)
        return ParsedDocument(source=str(path), parser="docx", blocks=blocks)
